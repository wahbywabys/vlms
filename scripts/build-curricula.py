#!/usr/bin/env python3
"""Turn the WABYS curriculum manual into LMS programmes.

    python3 scripts/build-curricula.py

Reads db/curricula/wabys-2026.json (every programme of the WABYS Manual for
Vocational Education & Applied Technology Curriculums, Revision 1.1, with its
knowledge modules, skill competencies and attitudes, in English and Arabic;
scripts/import-manual.py fills it from the HTML manual) and writes:

  db/migrations/007_wabys_rev11.sql             the seven programmes added in Revision 1.1, with their
                                                skills trees, and the Revision 1.1 profile of all fifteen
  app/public/templates/wabys-curricula-2026.xlsx  every programme in the programmes import format
  app/public/templates/wabys-curricula-2026.csv

Migration 006 (the first eight programmes, first edition) is already applied
in centres and is not regenerated; its trees stay as they are.

The skills tree follows the manual: for each programme and year there is one
unit per domain (Knowledge, Skills, Attitudes); each module, competency or
attitude is an element (its manual code, e.g. ELI-Y1-S-01); and its content
topics, performance criteria or observable behaviours are the criteria under
it (ELI-Y1-S-01.1, .2 …).
"""
import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'db' / 'curricula' / 'wabys-2026.json'
SQL = ROOT / 'db' / 'migrations' / '007_wabys_rev11.sql'
FIRST_EDITION = {'ELI', 'AUT', 'CNC', 'RMG', 'EIM', 'PRT', 'AIT', 'CWS'}
OUT = ROOT / 'app' / 'public' / 'templates'
EFFECTIVE = '2026-09-01'

data = json.loads(SRC.read_text(encoding='utf-8'))
YEARS, DOMAINS = data['years'], data['domains']
KSA = {'K': 'knowledge', 'S': 'skill', 'A': 'attitude'}

ICONS = {
    'rmg': 'M8 3 3 6l2 4 2-1v12h10V9l2 1 2-4-5-3a4 4 0 0 1-8 0z',
    'eim': 'M14.7 6.3a4 4 0 0 0-5.4 5.4L3 18l3 3 6.3-6.3a4 4 0 0 0 5.4-5.4l-2.5 2.5-2.5-.6-.6-2.5z',
    'prt': 'M6 9V3h12v6M6 18H4a2 2 0 0 1-2-2v-5a2 2 0 0 1 2-2h16a2 2 0 0 1 2 2v5a2 2 0 0 1-2 2h-2M6 14h12v7H6z',
    'ait': 'M8 8h8v8H8zM5 5h14v14H5zM9 2v3M15 2v3M9 19v3M15 19v3M2 9h3M2 15h3M19 9h3M19 15h3',
    'cws': 'M9 11a4 4 0 1 0 0-8 4 4 0 0 0 0 8zM2 21v-1a7 7 0 0 1 14 0v1M16 3.1a4 4 0 0 1 0 7.8M22 21v-1a7 7 0 0 0-4-6.3',
    'iel': 'M7 3v4M12 3v4M17 3v4M7 17v4M12 17v4M17 17v4M3 7h18v10H3zM8 12h8',
    'csn': 'M4 4h16v6H4zM4 14h16v6H4zM8 7h.01M8 17h.01M12 10v4',
    'rac': 'M12 2v20M4.9 4.9l14.2 14.2M2 12h20M4.9 19.1 19.1 4.9M9 3l3 3 3-3M9 21l3-3 3 3',
    'ham': 'M5 2h14v20H5zM5 7h14M12 15m-4 0a4 4 0 1 0 8 0 4 4 0 1 0-8 0M8 4.5h.01M11 4.5h.01',
    'oeq': 'M6 9V3h12v6M6 18H4a2 2 0 0 1-2-2v-5a2 2 0 0 1 2-2h16a2 2 0 0 1 2 2v5a2 2 0 0 1-2 2h-2M6 14h12v7H6zM17 12h.01',
    'ave': 'M2 5h20v12H2zM8 21h8M12 17v4M10 9l4 2-4 2z',
    'iem': 'M12 12m-6 0a6 6 0 1 0 12 0 6 6 0 1 0-12 0M12 12m-2 0a2 2 0 1 0 4 0 2 2 0 1 0-4 0M18 12h4M2 12h4',
}
# Skills whose failure in a practical is a safety gate (automatic "not yet competent", manual 1.3.2).
SAFETY = re.compile(r'\b(safe|safely|safety|LOTO|de-energi[sz]e|HV|SRS|lift)\b', re.I)


def q(v):
    if v is None:
        return 'NULL'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def jq(v):
    return q(json.dumps(v, ensure_ascii=False)) + '::jsonb'


def ar_join(parts):
    return '؛ '.join(parts)


rows_units, rows_elements, rows_criteria = [], [], []
sheet = []  # import rows
for p in data['programs']:
    fw = f"WABYS-{p['code']}-2026"
    trade = p['trade']
    items = p['items']
    in_sql = p['code'] not in FIRST_EDITION  # 006 already holds the first eight trees
    for y in (1, 2, 3):
        for d in 'KSA':
            its = [i for i in items if i['year'] == y and i['ksa'] == d]
            if not its:
                continue
            ucode = f"{p['code']}-Y{y}-{d}"
            yt = YEARS[str(y)]
            title = f"Year {y} · {DOMAINS[d]['en']}"
            title_ar = f"{yt['ar'].split(' · ')[0]} · {DOMAINS[d]['ar']}"
            hours = sum(i.get('hours') or 0 for i in its) or None
            in_sql and rows_units.append((fw, ucode, title, title_ar, yt['en'], yt['ar'], d, y, hours, 10 * y + 'KSA'.index(d)))
            sheet.append([p['code'], p['name']['en'], p['name']['ar'], ucode, 'unit', '', title, title_ar, d, y, '', '', '', '', hours or ''])
            for n, i in enumerate(its, 1):
                crit = i['criteria']
                safety = d == 'S' and bool(SAFETY.search(i['title']['en']))
                in_sql and rows_elements.append((fw, i['code'], ucode, i['title']['en'], i['title']['ar'],
                                      '; '.join(c['en'] for c in crit), ar_join([c['ar'] for c in crit]),
                                      d, y, i.get('hours'), i.get('level'), i.get('core'), safety, n))
                sheet.append([p['code'], '', '', i['code'], 'element', ucode, i['title']['en'], i['title']['ar'], d, y,
                              i.get('level') or '', i.get('core') or '', '; '.join(c['en'] for c in crit), ar_join([c['ar'] for c in crit]),
                              i.get('hours') or '', 'yes' if safety else ''])
                for m, c in enumerate(crit, 1):
                    in_sql and rows_criteria.append((fw, f"{i['code']}.{m}", i['code'], c['en'], c['ar'], d, y, m))
                    sheet.append([p['code'], '', '', f"{i['code']}.{m}", 'criterion', i['code'], c['en'], c['ar'], d, y, '', '', '', '', ''])


def values(rows):
    return ',\n'.join('  (' + ', '.join(q(v) for v in r) + ')' for r in rows)


programs = data['programs']
added = [p for p in programs if p['code'] not in FIRST_EDITION]
existing = {'cnc', 'auto', 'elec'}
out = []
w = out.append
w(f"""-- WABYS curricula, Revision 1.1 (generated by scripts/build-curricula.py from
-- db/curricula/wabys-2026.json; edit the JSON and run the script, do not edit by hand).
--
-- Source: {data['source']}
--
-- Revision 1.1 adds seven programmes from the PVTD occupations (IEL, CSN,
-- RAC, HAM, OEQ, AVE, IEM), the shared Electro-technology Foundation (ETF)
-- and Cross-Program Modules (CPM), and two study systems: A (60 %
-- competence standard) and B (80 % mastery standard, PVTD centres) with a
-- training-unit plan per new programme. The new programmes' trees include
-- the ETF and CPM units they adopt, under their shared codes, so a learner
-- who moves between electrical-family programmes keeps the credit.
-- The programme card (trades.profile) of all fifteen gains the Revision 1.1
-- fields: duration in both systems, PVTD reference occupation, shared units
-- and the System B plan.
SET search_path = vtc, public;
""")
for p in programs:
    profile = {k: p.get(k) for k in ('occupation', 'overview', 'eqf', 'hours', 'occupations', 'standards', 'equipment', 'tracks', 'capstone',
                                     'duration', 'pvtd', 'adopts', 'system_b')}
    profile['source'] = {'en': data['source'], 'ar': data['source_ar']}
    profile['revision'] = '1.1'
    if p['code'] in FIRST_EDITION:
        w(f"UPDATE trades SET profile = {jq(profile)}\n WHERE program_code = {q(p['code'])};")
        for r in p.get('revised', []):
            w(f"UPDATE competencies c SET title = {q(r['title']['en'])}, title_ar = {q(r['title']['ar'])}\n"
              f"  FROM competency_frameworks f WHERE f.id = c.framework_id AND f.code = 'WABYS-{p['code']}-2026' AND c.code = {q(r['code'])};")
    else:
        w(f"INSERT INTO trades (id, code, name, name_ar, icon, sort_order, program_code, accent, profile) VALUES ({q(p['trade'])}, {q(p['code'])}, "
          f"{q(p['name']['en'])}, {q(p['name']['ar'])}, {q(ICONS[p['trade']])}, {p['sort']}, {q(p['code'])}, {q(p['accent'])},\n"
          f"  {jq(profile)})\n  ON CONFLICT (id) DO UPDATE SET name = EXCLUDED.name, name_ar = EXCLUDED.name_ar, program_code = EXCLUDED.program_code,"
          f" accent = EXCLUDED.accent, profile = EXCLUDED.profile;")

w(f"""
INSERT INTO competency_frameworks (code, title, issuing_body, version, effective_from, trade_id) VALUES
{values([(f"WABYS-{p['code']}-2026", f"WABYS {p['name']['en']} (EQF {p['eqf'].replace('EQF ', '')})", 'WABYS for Training and Education', data['edition'], EFFECTIVE, p['trade']) for p in added])}
ON CONFLICT (code) DO UPDATE SET trade_id = EXCLUDED.trade_id, title = EXCLUDED.title, version = EXCLUDED.version;

-- -----------------------------------------------------------------------------
-- Skills tree: units (year × domain)
-- -----------------------------------------------------------------------------
INSERT INTO competencies (framework_id, code, title, title_ar, description, description_ar, kind, ksa, year, nominal_hours, sort_order)
SELECT f.id, v.code, v.title, v.title_ar, v.descr, v.descr_ar, 'unit', v.ksa, v.year, v.hours, v.sort
  FROM (VALUES
{values(rows_units)}
  ) AS v(fw, code, title, title_ar, descr, descr_ar, ksa, year, hours, sort)
  JOIN competency_frameworks f ON f.code = v.fw;

-- Elements: knowledge modules, skill competencies and attitudes (ETF-… and CPM-… are the shared units)
INSERT INTO competencies (framework_id, parent_id, code, title, title_ar, description, description_ar, kind, ksa, year, nominal_hours,
                          target_level, core_attitude, is_safety_critical, sort_order)
SELECT f.id, p.id, v.code, v.title, v.title_ar, v.descr, v.descr_ar, 'element', v.ksa, v.year, v.hours, v.lvl, v.core, v.safety, v.sort
  FROM (VALUES
{values(rows_elements)}
  ) AS v(fw, code, parent, title, title_ar, descr, descr_ar, ksa, year, hours, lvl, core, safety, sort)
  JOIN competency_frameworks f ON f.code = v.fw
  JOIN competencies p ON p.framework_id = f.id AND p.code = v.parent;

-- Criteria: content topics, performance criteria and observable behaviours
INSERT INTO competencies (framework_id, parent_id, code, title, title_ar, kind, ksa, year, sort_order)
SELECT f.id, p.id, v.code, v.title, v.title_ar, 'criterion', v.ksa, v.year, v.sort
  FROM (VALUES
{values(rows_criteria)}
  ) AS v(fw, code, parent, title, title_ar, ksa, year, sort)
  JOIN competency_frameworks f ON f.code = v.fw
  JOIN competencies p ON p.framework_id = f.id AND p.code = v.parent;

-- -----------------------------------------------------------------------------
-- Machine types for the new programmes and their 3D twins
-- -----------------------------------------------------------------------------
INSERT INTO machine_types (id, trade_id, name, name_ar, model_id, twin_match, precheck_code, sort_order) VALUES
  ('electronics_bench', 'iel', 'Electronics bench (scope, generator, supply, rework)', 'طاولة إلكترونيات (راسم موجات ومولد ومصدر ولحام)', 'electronics-bench', 'exact', NULL, 90),
  ('plc_trainer',       'iel', 'PLC and process-control trainer',    'وحدة تدريب PLC والتحكم في العمليات',   'electronics-bench', 'similar', NULL, 91),
  ('pcb_line',          'iel', 'PCB exposure and etching line',      'خط تعريض وحفر الدوائر المطبوعة',       NULL,                NULL,      NULL, 92),
  ('network_rack',      'csn', 'Network rack (switch, router, server, UPS)', 'خزانة شبكات (سويتش وموجه وخادم وUPS)', 'network-rack', 'exact', NULL, 100),
  ('pc_bench',          'csn', 'PC assembly bench',                  'طاولة تجميع حاسبات',                   'network-rack',      'similar', NULL, 101),
  ('split_ac',          'rac', 'Split air-conditioner service rig',  'منصة صيانة مكيف سبليت',                 'split-ac',          'exact',   NULL, 110),
  ('fridge_trainer',    'rac', 'Domestic refrigeration trainer',     'وحدة تدريب التبريد المنزلي',           'split-ac',          'similar', NULL, 111),
  ('recovery_unit',     'rac', 'Refrigerant recovery machine',       'ماكينة استرجاع وسيط التبريد',          NULL,                NULL,      NULL, 112),
  ('washing_machine',   'ham', 'Washing machine on repair bench',    'غسالة على طاولة الإصلاح',              'washing-machine',   'exact',   NULL, 120),
  ('dishwasher',        'ham', 'Dishwasher trainer',                 'وحدة تدريب غسالة أطباق',               'washing-machine',   'similar', NULL, 121),
  ('pat_tester',        'ham', 'Appliance safety tester (PAT)',      'جهاز اختبار سلامة الأجهزة',            NULL,                NULL,      NULL, 122),
  ('laser_mfp',         'oeq', 'Laser multifunction printer (MFP)',  'طابعة ليزر متعددة الوظائف',            'laser-mfp',         'exact',   NULL, 130),
  ('copier',            'oeq', 'Photocopier',                        'ماكينة تصوير مستندات',                 'laser-mfp',         'similar', NULL, 131),
  ('pos_terminal',      'oeq', 'POS terminal and banknote counter',  'جهاز نقاط بيع وعداد نقود',             NULL,                NULL,      NULL, 132),
  ('tv_bench',          'ave', 'TV / display repair bench',          'طاولة إصلاح التلفزيون والشاشات',       'tv-bench',          'exact',   NULL, 140),
  ('cctv_kit',          'ave', 'IP-CCTV kit with NVR',               'طقم كاميرات IP مع جهاز تسجيل',          NULL,                NULL,      NULL, 141),
  ('rewind_bench',      'iem', 'Motor rewinding and test bench',     'طاولة لف واختبار المحركات',            'motor-rewind',      'exact',   NULL, 150),
  ('motor_control',     'iem', 'Motor-control bench (DOL, star-delta, VFD)', 'طاولة التحكم في المحركات',       'dol-panel',         'similar', NULL, 151),
  ('transformer_trainer','iem', 'Transformer trainer',                'وحدة تدريب المحولات',                  NULL,                NULL,      NULL, 152)
ON CONFLICT (id) DO NOTHING;
""")
SQL.write_text('\n'.join(out), encoding='utf-8')

# ---------------------------------------------------------------------------
# The same rows as an import workbook (re-importable from the programmes screen)
# ---------------------------------------------------------------------------
COLS = [  # English, Arabic, width
    ('Programme code', 'رمز التخصص', 12), ('Programme name', 'اسم التخصص', 26), ('Programme name (Arabic)', 'اسم التخصص بالعربية', 26),
    ('Skill code', 'رمز المهارة', 18), ('Level', 'المستوى', 10), ('Parent code', 'رمز المستوى الأعلى', 16),
    ('Skill title', 'عنوان المهارة', 44), ('Skill title (Arabic)', 'عنوان المهارة بالعربية', 44),
    ('Domain (K/S/A)', 'المجال (معرفة/مهارة/اتجاه)', 10), ('Year', 'السنة', 7), ('Target level', 'مستوى الأداء المستهدف', 10),
    ('Core attitude', 'الاتجاه الأساسي', 10), ('Description', 'الوصف', 60), ('Description (Arabic)', 'الوصف بالعربية', 60),
    ('Hours', 'الساعات', 8), ('Safety-critical', 'سلامة حرجة', 10),
]
thin = Side(style='thin', color='C9CED8')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {'unit': 'DCE3F2', 'element': 'FFFFFF', 'criterion': 'F7F8FA'}
wb = Workbook()
ws = wb.active
ws.title = 'Programmes'
for i, (en, ar, width) in enumerate(COLS, 1):
    c = ws.cell(row=1, column=i, value=f'{en}\n{ar}')
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='14234C')
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.column_dimensions[c.column_letter].width = width
ws.row_dimensions[1].height = 42
ws.freeze_panes = 'E2'
for r, row in enumerate(sheet, 2):
    row = row + [''] * (len(COLS) - len(row))
    for i, v in enumerate(row, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = BOX
        c.fill = PatternFill('solid', fgColor=FILL[row[4]])
        c.alignment = Alignment(wrap_text=i in (7, 8, 13, 14), vertical='top')
        if row[4] == 'unit':
            c.font = Font(bold=True, color='14234C')
att = wb.create_sheet('Core attitudes')
att.append(['Code', 'Dimension', 'البعد', 'Observable indicators', 'المؤشرات الملاحظة'])
for a in data['attitudes']:
    att.append([a['code'], a['en'], a['ar'], a['ind_en'], a['ind_ar']])
for col, wd in zip('ABCDE', (8, 32, 32, 60, 60)):
    att.column_dimensions[col].width = wd
prog = wb.create_sheet('Programme profiles')
prog.append(['Code', 'Programme', 'التخصص', 'EQF', 'Hours', 'Occupation', 'المهنة', 'Overview', 'نظرة عامة', 'Capstone', 'مشروع التخرج'])
for p in programs:
    prog.append([p['code'], p['name']['en'], p['name']['ar'], p['eqf'], p['hours'], p['occupation']['en'], p['occupation']['ar'],
                 p['overview']['en'], p['overview']['ar'], p['capstone']['en'], p['capstone']['ar']])
for col in 'HIJK':
    prog.column_dimensions[col].width = 70
OUT.mkdir(parents=True, exist_ok=True)
wb.save(OUT / 'wabys-curricula-2026.xlsx')
with (OUT / 'wabys-curricula-2026.csv').open('w', encoding='utf-8-sig', newline='') as f:
    cw = csv.writer(f)
    cw.writerow([c[0] for c in COLS])
    for row in sheet:
        cw.writerow(row + [''] * (len(COLS) - len(row)))

print(f'{len(programs)} programmes, {len(rows_units)} units, {len(rows_elements)} competencies, {len(rows_criteria)} criteria')
print(f'wrote {SQL.relative_to(ROOT)} and app/public/templates/wabys-curricula-2026.xlsx/.csv')
