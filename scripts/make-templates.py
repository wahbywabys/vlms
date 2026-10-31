#!/usr/bin/env python3
"""Build the Excel templates the admin screens import.

    python3 scripts/make-templates.py        # writes app/public/templates/*.xlsx and *.csv

  machine-inventory-template.xlsx  one row per real machine in each centre's workshops
  programmes-template.xlsx         programmes (specialisations) and their skills tree

The machine types (and the 3D twin each maps to) are read from
db/migrations/004_centers_programs_machines.sql and 006_wabys_curricula.sql, so the dropdown always
matches what the import accepts. Column headers are "English\nالعربية";
the import reads either line, so a translated or reordered sheet still works.
"""
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'app' / 'public' / 'templates'
NAVY, ORANGE, SOFT, GREY = '14234C', 'F06721', 'FDEBDD', 'F3F4F7'
TRADES = {'cnc': 'CNC machining', 'auto': 'Automotive', 'elec': 'Electrical', 'weld': 'Welding', 'rmg': 'Garments',
          'eim': 'Industrial maintenance', 'prt': 'Printing', 'ait': 'AI for technicians', 'cws': 'Work skills', None: 'Any'}

TYPES = []
for mig in ['004_centers_programs_machines.sql', '006_wabys_curricula.sql']:
    sql = (ROOT / 'db' / 'migrations' / mig).read_text(encoding='utf-8')
    for m in re.finditer(r"\('([a-z_]+)',\s+(NULL|'[a-z]+'),\s+'([^']+)',\s+'([^']+)',\s+(NULL|'[a-z-]+'),\s+(NULL|'[a-z]+'),", sql):
        tid, trade, en, ar, model, match = m.groups()
        unq = lambda v: None if v == 'NULL' else v.strip("'")
        TYPES.append({'id': tid, 'trade': unq(trade), 'en': en, 'ar': ar, 'model': unq(model), 'match': unq(match)})
    # Later migrations point existing types at new 3D twins.
    for m in re.finditer(r"UPDATE machine_types SET model_id = '([a-z-]+)', twin_match = '([a-z]+)' WHERE id = '([a-z_]+)'", sql):
        for t in TYPES:
            if t['id'] == m.group(3):
                t['model'], t['match'] = m.group(1), m.group(2)
assert len(TYPES) >= 24, 'machine types not found in the migrations'

thin = Side(style='thin', color='C9CED8')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, cols):
    for i, (en, ar, width, required, _) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=f"{en}{' *' if required else ''}\n{ar}")
        c.font = Font(bold=True, color='FFFFFF' if required else NAVY)
        c.fill = PatternFill('solid', fgColor=ORANGE if required else 'DCE3F2')
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = 'A2'


def body_style(ws, cols, rows=500):
    for r in range(2, rows + 2):
        for i in range(1, len(cols) + 1):
            ws.cell(row=r, column=i).border = BOX


def instructions(ws, title, lines):
    ws.sheet_view.rightToLeft = False
    ws.column_dimensions['A'].width = 110
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=16, color=NAVY)
    for i, (text, bold) in enumerate(lines, 3):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, color=NAVY if bold else '222222')
        c.alignment = Alignment(wrap_text=True, vertical='top')


def lists_sheet(wb, name, columns):
    ws = wb.create_sheet(name)
    for ci, values in enumerate(columns, 1):
        for ri, v in enumerate(values, 1):
            ws.cell(row=ri, column=ci, value=v)
    ws.sheet_state = 'hidden'
    return ws


def dropdown(ws, col_letter, ref, rows=500, prompt=None):
    dv = DataValidation(type='list', formula1=ref, allow_blank=True, showErrorMessage=True,
                        errorTitle='Pick from the list', error='Choose a value from the dropdown list.')
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = 'Tip', prompt, True
    ws.add_data_validation(dv)
    dv.add(f'{col_letter}2:{col_letter}{rows + 1}')


def write_csv(path, cols, examples):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in cols])
        for row in examples:
            w.writerow(row)


# ---------------------------------------------------------------------------
# Machine inventory
# ---------------------------------------------------------------------------
M_COLS = [  # (English, Arabic, width, required, import key)
    ('Center code', 'رمز المركز', 14, True, 'centerCode'),
    ('Center name', 'اسم المركز', 26, False, 'centerName'),
    ('Center name (Arabic)', 'اسم المركز بالعربية', 26, False, 'centerNameAr'),
    ('Governorate', 'المحافظة', 14, False, 'region'),
    ('Workshop', 'الورشة', 24, True, 'workshop'),
    ('Workshop (Arabic)', 'الورشة بالعربية', 24, False, 'workshopAr'),
    ('Machine type', 'نوع الماكينة', 30, True, 'machineType'),
    ('Brand', 'الماركة', 16, True, 'brand'),
    ('Model', 'الموديل', 18, True, 'model'),
    ('Serial number', 'الرقم التسلسلي', 18, True, 'serial'),
    ('Asset tag (leave blank)', 'رقم الأصل (اتركه فارغًا)', 18, False, 'assetTag'),
    ('Year made', 'سنة الصنع', 10, False, 'year'),
    ('Status', 'الحالة', 16, False, 'status'),
    ('Required skill code', 'رمز المهارة المطلوبة', 18, False, 'competencyCode'),
    ('Photo file name', 'اسم ملف الصورة', 22, False, 'photo'),
    ('Notes', 'ملاحظات', 30, False, 'notes'),
]
M_EXAMPLES = [
    ['ALX-01', 'VTC Alexandria', 'مركز التدريب المهني بالإسكندرية', 'Alexandria', 'Turning shop', 'ورشة الخراطة', 'CNC lathe', 'Emco', 'Concept Turn 250', 'EM-2201-0457', '', 2019, 'In service', '', 'ALX-lathe-1.jpg', ''],
    ['ALX-01', 'VTC Alexandria', 'مركز التدريب المهني بالإسكندرية', 'Alexandria', 'Turning shop', 'ورشة الخراطة', 'Centre lathe (manual)', 'Harrison', 'M300', 'H300-1188', '', 2008, 'Maintenance', '', '', 'Tailstock being repaired'],
    ['ALX-01', 'VTC Alexandria', 'مركز التدريب المهني بالإسكندرية', 'Alexandria', 'Welding shop', 'ورشة اللحام', 'MIG/MAG welder', 'Lincoln', 'Powertec i250C', 'U1190312345', '', 2021, 'In service', '', 'ALX-mig-1.jpg', ''],
]
STATUSES = ['In service', 'Maintenance', 'Out of service']

wb = Workbook()
ins = wb.active
ins.title = 'Instructions'
instructions(ins, 'VTC_LMS · Machine inventory template · نموذج حصر الماكينات', [
    ('What this is for', True),
    ('List every real machine in every workshop of your centre, one row per machine. When you import this file in VTC_LMS '
     '(Director → Machine register → Import), each machine gets its own record, its own QR sticker, and is linked to the 3D '
     'model of its type so learners practise on the matching twin before they touch the real one.', False),
    ('الغرض من النموذج', True),
    ('اكتب كل ماكينة حقيقية في كل ورشة بالمركز، ماكينة واحدة في كل سطر. عند استيراد الملف في VTC_LMS (المدير ← سجل الماكينات ← استيراد) '
     'تحصل كل ماكينة على سجل خاص وملصق QR خاص، وتُربط بالنموذج ثلاثي الأبعاد لنوعها ليتدرب المتدرب على نظيرها الافتراضي قبل الحقيقية.', False),
    ('How to fill it in', True),
    ('1. Fill the "Machines" sheet. Orange columns are required. Use the dropdowns for Machine type and Status.', False),
    ('2. Center code is short and fixed for each centre (for example ALX-01). Use the same code on every row of that centre.', False),
    ('3. Serial number is on the machine nameplate. It is how the LMS recognises the same machine if you import again.', False),
    ('4. Leave "Asset tag" blank: the LMS makes one (for example ALX-01-CNCLAT-01) and prints it on the sticker.', False),
    ('5. Photo: take one clear photo of each machine with its nameplate visible, name the file, and write the file name here. '
     'Select the photos together with this file when you import.', False),
    ('6. If a type is missing from the list, choose "Other machine" and describe it in Notes. It is registered without a 3D twin.', False),
    ('7. Required skill code (optional): the code of a skill from Programmes that a learner must hold before using this machine.', False),
    ('طريقة التعبئة', True),
    ('١. املأ ورقة "Machines". الأعمدة البرتقالية إلزامية. استخدم القوائم المنسدلة لنوع الماكينة والحالة.', False),
    ('٢. رمز المركز قصير وثابت لكل مركز (مثل ALX-01)، ويُكتب بنفس الشكل في كل سطور المركز.', False),
    ('٣. الرقم التسلسلي مكتوب على لوحة بيانات الماكينة، وبه يتعرف النظام على نفس الماكينة إذا استوردت الملف مرة أخرى.', False),
    ('٤. اترك "رقم الأصل" فارغًا؛ النظام ينشئه ويطبعه على الملصق.', False),
    ('٥. الصورة: صوّر كل ماكينة صورة واضحة تظهر فيها لوحة البيانات، واكتب اسم الملف هنا، واختر الصور مع هذا الملف عند الاستيراد.', False),
    ('٦. إذا لم تجد النوع في القائمة اختر "Other machine" واكتب وصفها في الملاحظات؛ تُسجّل بدون نموذج ثلاثي الأبعاد.', False),
    ('Examples (do not copy into the Machines sheet as they are)', True),
    *[(' | '.join(str(v) for v in row if v != ''), False) for row in M_EXAMPLES],
])

ws = wb.create_sheet('Machines')
header(ws, M_COLS)
body_style(ws, M_COLS)
ref = wb.create_sheet('Machine types')
ref_cols = [('Machine type (pick this)', 'نوع الماكينة', 30, True, ''), ('Arabic name', 'الاسم بالعربية', 30, False, ''),
            ('Programme', 'التخصص', 16, False, ''), ('3D twin in VTC_LMS', 'النموذج ثلاثي الأبعاد', 22, False, ''),
            ('Match', 'التطابق', 30, False, '')]
header(ref, ref_cols)
for r, t in enumerate(TYPES, 2):
    match = {'exact': 'Same machine type', 'similar': 'Closest twin (similar controls and hazards)', None: 'No 3D twin yet'}[t['match']]
    for c, v in enumerate([t['en'], t['ar'], TRADES[t['trade']], t['model'] or '—', match], 1):
        ref.cell(row=r, column=c, value=v).border = BOX
lists_sheet(wb, 'Lists', [[t['en'] for t in TYPES], STATUSES])
dropdown(ws, 'G', f"=Lists!$A$1:$A${len(TYPES)}", prompt='Pick the machine type. See the "Machine types" sheet for the 3D twin of each.')
dropdown(ws, 'M', f"=Lists!$B$1:$B${len(STATUSES)}")
wb.active = 1
OUT.mkdir(parents=True, exist_ok=True)
wb.save(OUT / 'machine-inventory-template.xlsx')
write_csv(OUT / 'machine-inventory-template.csv', M_COLS, M_EXAMPLES)

# ---------------------------------------------------------------------------
# Programmes and skills
# ---------------------------------------------------------------------------
P_COLS = [
    ('Programme code', 'رمز التخصص', 16, True, 'programCode'),
    ('Programme name', 'اسم التخصص', 26, False, 'programName'),
    ('Programme name (Arabic)', 'اسم التخصص بالعربية', 26, False, 'programNameAr'),
    ('Skill code', 'رمز المهارة', 20, False, 'competencyCode'),
    ('Level', 'المستوى', 12, False, 'level'),
    ('Parent code', 'رمز المستوى الأعلى', 20, False, 'parentCode'),
    ('Skill title', 'عنوان المهارة', 44, False, 'title'),
    ('Skill title (Arabic)', 'عنوان المهارة بالعربية', 44, False, 'titleAr'),
    ('Safety-critical', 'سلامة حرجة', 12, False, 'safetyCritical'),
    ('Hours', 'الساعات', 10, False, 'hours'),
    ('Domain (K/S/A)', 'المجال (معرفة/مهارة/اتجاه)', 12, False, 'ksa'),
    ('Year', 'السنة', 8, False, 'year'),
    ('Target level', 'مستوى الأداء المستهدف', 12, False, 'targetLevel'),
    ('Core attitude', 'الاتجاه الأساسي', 12, False, 'coreAttitude'),
]
P_EXAMPLES = [
    ['PLMB', 'Plumbing', 'السباكة', '', '', '', '', '', '', '', '', '', '', ''],
    ['PLMB', '', '', 'PLMB-Y1-S', 'Unit', '', 'Year 1 · Practical skills', 'السنة الأولى · المهارات العملية', 'No', '', 'S', 1, '', ''],
    ['PLMB', '', '', 'PLMB-Y1-S-01', 'Element', 'PLMB-Y1-S', 'Cut and join copper pipe', 'قطع ووصل المواسير النحاس', 'No', 12, 'S', 1, 'L2', ''],
    ['PLMB', '', '', 'PLMB-Y1-S-01.1', 'Criterion', 'PLMB-Y1-S-01', 'Uses the blowtorch safely', 'يستخدم لمبة اللحام بأمان', 'Yes', '', 'S', 1, '', ''],
]
wb = Workbook()
ins = wb.active
ins.title = 'Instructions'
instructions(ins, 'VTC_LMS · Programmes and skills template · نموذج التخصصات والمهارات', [
    ('How it works', True),
    ('Each programme (specialisation) has a tree of skills: Unit → Element → Criterion. One row per skill. A row with only the '
     'programme columns just creates or renames the programme. Import it in VTC_LMS: Director → Programmes & skills → Import.', False),
    ('Write each parent above its children: a Unit row before its Elements, an Element row before its Criteria. Parent code is '
     'empty for a Unit. Mark a skill Safety-critical "Yes" when failing it must stop the learner (it then gates machine access).', False),
    ('Importing the same file again updates titles and hours in place; nothing is duplicated.', False),
    ('Optional, as in the WABYS curriculum manual: Domain K (knowledge), S (skill) or A (attitude); Year 1–3; Target level L1–L4 for '
     'skills; Core attitude A1–A7 for attitudes. The full manual is in wabys-curricula-2026.xlsx.', False),
    ('طريقة العمل', True),
    ('لكل تخصص شجرة مهارات: وحدة ← عنصر ← معيار أداء. سطر واحد لكل مهارة. السطر الذي يحتوي أعمدة التخصص فقط ينشئ التخصص أو يعدّل اسمه. '
     'اكتب المستوى الأعلى قبل ما تحته، واترك "رمز المستوى الأعلى" فارغًا للوحدة. اختر "Yes" في السلامة الحرجة للمهارة التي يمنع الرسوب فيها المتدرب من الاستمرار.', False),
    ('Example', True),
    *[(' | '.join(str(v) for v in row if v != ''), False) for row in P_EXAMPLES],
])
ws = wb.create_sheet('Programmes')
header(ws, P_COLS)
body_style(ws, P_COLS)
lists_sheet(wb, 'Lists', [['Unit', 'Element', 'Criterion'], ['Yes', 'No'], ['K', 'S', 'A'], ['L1', 'L2', 'L3', 'L4']])
dropdown(ws, 'E', '=Lists!$A$1:$A$3')
dropdown(ws, 'I', '=Lists!$B$1:$B$2')
dropdown(ws, 'K', '=Lists!$C$1:$C$3')
dropdown(ws, 'M', '=Lists!$D$1:$D$4')
wb.active = 1
wb.save(OUT / 'programmes-template.xlsx')
write_csv(OUT / 'programmes-template.csv', P_COLS, P_EXAMPLES)
print(f'wrote {len(TYPES)} machine types into {OUT}')
